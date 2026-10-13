"""Native XLSX evidence. This module never imports a PageIndex or PDF parser."""
import hashlib
import io
import json
from datetime import date, datetime, time
from zipfile import ZipFile, BadZipFile

import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter

PARSER_VERSION = f'xlsx-cells-1:openpyxl-{openpyxl.__version__}'


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def scalar(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or type(value) in (str, bool, int, float):
        return value
    raise ValueError('Unsupported spreadsheet value type')


def bounds(value, limit=10000):
    try:
        result = range_boundaries(value)
        if any(type(v) is not int or v < 1 for v in result):
            raise ValueError()
        if result[0] > result[2] or result[1] > result[3]:
            raise ValueError()
        if (result[2] - result[0] + 1) * (result[3] - result[1] + 1) > limit:
            raise ValueError()
        return result
    except (ValueError, TypeError):
        raise ValueError('Invalid or oversized cell range') from None


def parse(data, configured_tables=None):
    configured_tables = configured_tables or []
    for item in configured_tables:
        if not isinstance(item, dict) or set(item) != {'id', 'sheet', 'range', 'header_row'} or any(not isinstance(item[key], str) or not item[key] for key in ['id', 'sheet', 'range']) or type(item['header_row']) is not int:
            raise ValueError('Table configuration requires explicit id, sheet, range and header_row')
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if 'xl/workbook.xml' not in archive.namelist():
                raise ValueError('Source is not an XLSX workbook')
            if sum(info.file_size for info in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError('Workbook exceeds uncompressed size limit')
    except BadZipFile:
        raise ValueError('Source is not an XLSX workbook') from None
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    try:
        sheets = []
        for index, sheet in enumerate(formulas.worksheets, 1):
            ident = hashlib.sha256(encode([index, sheet.title]).encode()).hexdigest()[:24]
            cells = {}
            if sheet.max_row * sheet.max_column > 1000000:
                raise ValueError('Workbook sheet exceeds cell limit')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    formula = cell.value if cell.data_type == 'f' else None
                    value = cached[sheet.title][cell.coordinate].value if formula else cell.value
                    cells[cell.coordinate] = {'coordinate': cell.coordinate, 'row': cell.row,
                        'column': cell.column, 'value': scalar(value), 'type': type(value).__name__,
                        'number_format': cell.number_format, 'formula': formula,
                        'cached_value_available': not formula or value is not None}
            tables = []
            definitions = [{'id': name, 'range': table.ref,
                            'header_row': range_boundaries(table.ref)[1] if table.headerRowCount else None,
                            'totals_rows': table.totalsRowCount or 0}
                           for name in sheet.tables for table in [sheet.tables[name]]]
            definitions += [entry for entry in (configured_tables or []) if entry.get('sheet') == sheet.title]
            for entry in definitions:
                left, top, right, bottom = bounds(entry['range'], 1000000)
                header = entry.get('header_row')
                if header is None:
                    continue  # Raw range reading remains available; never guess a header.
                if type(header) is not int or not top <= header <= bottom:
                    raise ValueError('Configured header is outside its table')
                if right > sheet.max_column or bottom > sheet.max_row: raise ValueError('Table outside worksheet')
                if any(item['id'] == str(entry['id']) for item in tables): raise ValueError('Duplicate table ID')
                column_types = {}
                for cell in cells.values():
                    if left <= cell['column'] <= right and header < cell['row'] <= bottom:
                        column_types.setdefault(cell['column'], set()).add(cell['type'])
                columns = []
                for column in range(left, right + 1):
                    key = get_column_letter(column)
                    columns.append({'id': key, 'header': scalar(sheet.cell(header, column).value),
                        'coordinate': f'{key}{header}', 'types': sorted(column_types.get(column, set()))})
                tables.append({'id': str(entry['id']), 'range': entry['range'], 'header_row': header,
                    'first_data_row': header + 1, 'last_data_row': bottom - entry.get('totals_rows', 0),
                    'columns': columns})
            sheets.append({'id': ident, 'name': sheet.title, 'index': index, 'visibility': sheet.sheet_state,
                'max_row': sheet.max_row, 'max_column': sheet.max_column,
                'merged_ranges': [{'range': str(area), 'anchor': f'{get_column_letter(area.min_col)}{area.min_row}'}
                                  for area in sheet.merged_cells.ranges], 'tables': tables, 'cells': cells})
        digest = hashlib.sha256(data).hexdigest()
        version = hashlib.sha256(encode([digest, PARSER_VERSION, configured_tables or [], sheets]).encode()).hexdigest()
        return {'sha256': digest, 'version': version, 'parser_version': PARSER_VERSION, 'sheets': sheets}
    finally:
        formulas.close(); cached.close()


def select(record, sheet_id, cell_range=None, table_id=None, filters=None, *, row_start=0, row_limit=None):
    sheet = next((item for item in record['sheets'] if item['id'] == sheet_id), None)
    if sheet is None:
        raise ValueError('Unknown worksheet ID')
    if bool(cell_range) == bool(table_id):
        raise ValueError('Specify exactly one range or table_id')
    if cell_range and filters:
        raise ValueError('Filters require a discovered table')
    if table_id:
        table = next((item for item in sheet['tables'] if item['id'] == table_id), None)
        if table is None:
            raise ValueError('Unknown table ID')
        left, _, right, _ = bounds(table['range'], 1000000)
        top, bottom = table['first_data_row'], table['last_data_row']
        known = {column['id'] for column in table['columns']}
        for condition in filters or []:
            if condition['column_id'] not in known or condition['operator'] not in ('eq', 'contains'):
                raise ValueError('Invalid column or filter operator')
    else:
        left, top, right, bottom = bounds(cell_range)
    if right > sheet['max_column'] or bottom > sheet['max_row']:
        raise ValueError('Cell range outside worksheet')
    result = []
    stop = max(top, row_start) - 1
    for row in range(max(top, row_start), bottom + 1):
        stop = row
        def matches(condition):
            cell = sheet['cells'].get(f"{condition['column_id']}{row}", {})
            if cell.get('cached_value_available') is False: raise ValueError('Cannot filter on an unavailable formula result')
            value = cell.get('value')
            expected = condition['value']
            if condition['operator'] == 'contains':
                return type(value) is str and type(expected) is str and expected in value
            # XLSX numeric cells can round-trip as either int or float. Equal
            # numbers retain their original returned type; strings and bools
            # are never coerced into the numeric domain.
            same_domain = (type(value) in (int, float) and type(expected) in (int, float)) or type(value) is type(expected)
            return same_domain and value == expected
        if not all(matches(condition) for condition in filters or []):
            continue
        area = f'{get_column_letter(left)}{row}:{get_column_letter(right)}{row}'
        selected = [sheet['cells'].get(f'{get_column_letter(column)}{row}', {
            'coordinate': f'{get_column_letter(column)}{row}', 'row': row, 'column': column,
            'value': None, 'type': 'NoneType', 'formula': None, 'cached_value_available': True})
            for column in range(left, right + 1)]
        locator = {'kind': 'cells', 'sheet_id': sheet['id'], 'sheet': sheet['name'], 'range': area}
        evidence_id = hashlib.sha256(encode([record['version'], locator]).encode()).hexdigest()
        result.append({'id': evidence_id, 'locator': locator, 'text': encode(selected)})
        if row_limit and len(result) >= row_limit: break
    return sheet | {'next_row': stop + 1 if stop < bottom else None}, result
