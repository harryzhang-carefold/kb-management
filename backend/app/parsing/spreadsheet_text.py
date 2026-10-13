"""Row-aware discovery text. Exact business evidence still comes from native cells.

Only XLSX tables or explicit editor configuration establish headers. No inferred
header, cross-row inheritance or business-code mapping is introduced here.
"""
import hashlib

from openpyxl.utils.cell import get_column_letter
from app.knowledge.spreadsheet import bounds, encode

MAX_BATCH_CHARS = 32000
LAYOUT_VERSION = 'xlsx-text-rows-1'


def _value(cell):
    if cell['formula'] is None:
        return cell['value']
    return {'formula': cell['formula'], 'cached_value': cell['value'],
            'cached_value_available': cell['cached_value_available']}


def _regions(sheet):
    # A cell can explicitly occur in overlapping configured tables. Such tables
    # are kept separate; their relationships are not inferred by the parser.
    areas = []
    for table in sheet['tables']:
        left, top, right, bottom = bounds(table['range'], 1000000)
        selected = [c for c in sheet['cells'].values()
                    if left <= c['column'] <= right and top <= c['row'] <= bottom and c['row'] != table['header_row']]
        areas.append((table, selected))
    ranges = [bounds(t['range'], 1000000) for t in sheet['tables']]
    raw = [c for c in sheet['cells'].values() if not any(
        left <= c['column'] <= right and top <= c['row'] <= bottom for left, top, right, bottom in ranges)]
    if raw:
        left = min(c['column'] for c in raw); right = max(c['column'] for c in raw)
        top = min(c['row'] for c in raw); bottom = max(c['row'] for c in raw)
        areas.append(({'id': None, 'range': f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}',
                       'header_row': None, 'columns': [], 'raw_region': True}, raw))
    return areas


def _fragments(row, context, limit):
    source = encode(row)
    if len(context) + len(source) + 1 <= limit:
        return [source]
    fragments = []; start = 0
    while start < len(source):
        def entry(end):
            return {'row': row['row'], 'range': row['range'], 'row_complete': False,
                    'continuation': {'part': len(fragments) + 1, 'parts': None,
                        'start': start, 'end': end, 'total_chars': len(source)},
                    'row_fragment': source[start:end]}
        low, high = start, len(source)
        # Reserve room for the final part count. JSON escaping is measured, not
        # estimated, so quotes/newlines/Unicode never bypass the declared bound.
        while low < high:
            middle = (low + high + 1) // 2
            if len(context) + len(encode(entry(middle))) + 33 <= limit:
                low = middle
            else:
                high = middle - 1
        if low == start:
            raise ValueError('Spreadsheet context leaves no room for a row fragment')
        fragments.append(entry(low)); start = low
    for fragment in fragments:
        fragment['continuation']['parts'] = len(fragments)
    return [encode(fragment) for fragment in fragments]


def render(record, limit=MAX_BATCH_CHARS):
    """Publish bounded, contiguous text batches with repeated source context."""
    parts = []
    for sheet in record['sheets']:
        for table, cells in _regions(sheet):
            context = ('## worksheet=' + encode({'id': sheet['id'], 'name': sheet['name']}) + '\n'
                       + '### region=' + encode({k: v for k, v in table.items() if k != 'columns'}) + '\n'
                       + 'columns=' + encode(table['columns']) + '\n'
                       + 'Row records are independent unless the source explicitly defines a cross-row relationship.\n'
                       + 'Missing coordinates inside a row range are blank cells, not inherited values.\n')
            if len(context) + 1 >= limit:
                raise ValueError('Spreadsheet header/context exceeds the text batch bound')
            by_row = {}
            for cell in sorted(cells, key=lambda c: (c['row'], c['column'])):
                by_row.setdefault(cell['row'], []).append(cell)
            batch = context
            for number, row_cells in by_row.items():
                left, _, right, _ = bounds(table['range'], 1000000)
                row = {'row': number, 'range': f'{get_column_letter(left)}{number}:{get_column_letter(right)}{number}',
                       'row_complete': True, 'cells': {c['coordinate']: _value(c) for c in row_cells}}
                for fragment in _fragments(row, context, limit):
                    line = fragment + '\n'
                    if len(batch) + len(line) > limit:
                        parts.append(batch); batch = context
                    batch += line
            parts.append(batch)
    if not parts:
        raise ValueError('XLSX 未提取到任何单元格内容')
    text = ''.join(parts); position = 0; batches = []
    for part in parts:
        batches.append([position, position + len(part)]); position += len(part)
    return text, {'spreadsheet_text': {'layout_version': LAYOUT_VERSION,
        'text_sha256': hashlib.sha256(text.encode()).hexdigest(), 'limit': limit, 'batches': batches}}


def partitions(text, layout, limit=MAX_BATCH_CHARS):
    """Validate parser-owned offsets before a builder uses them."""
    if layout.get('layout_version') != LAYOUT_VERSION or layout.get('text_sha256') != hashlib.sha256(text.encode()).hexdigest():
        raise ValueError('Spreadsheet text layout does not match the parsed source')
    pieces = []; position = 0
    for start, end in layout.get('batches', []):
        if type(start) is not int or type(end) is not int or start != position or not start < end <= len(text) or end - start > limit:
            raise ValueError('Invalid spreadsheet text batch boundary')
        pieces.append((start, text[start:end])); position = end
    if position != len(text) or not pieces:
        raise ValueError('Spreadsheet text layout does not cover the source')
    return pieces
