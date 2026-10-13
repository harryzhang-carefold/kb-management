import io
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.knowledge import spreadsheet
from app.knowledge.reading import read_batch, fully_read, ReadingCursorError


def workbook():
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Rules'
    sheet.append(['Code Type', 'Code', 'Action'])
    sheet.append(['Reason', '197', 'Check the actual authorization.'])
    sheet.append(['Status', '197', 'Check the actual status.'])
    sheet.append(['Reason', 197, 'Numeric identifier'])
    sheet.add_table(Table(displayName='RulesTable', ref='A1:C4'))
    other = book.create_sheet('Hidden'); other.sheet_state = 'hidden'
    other['A1'] = '=1+1'; other.merge_cells('B2:C2'); other['B2'] = 'Merged'
    stream = io.BytesIO(); book.save(stream); return stream.getvalue()


def test_native_workbook_preserves_types_formulas_merges_and_real_coordinates():
    record = spreadsheet.parse(workbook()); sheet = record['sheets'][0]
    assert sheet['tables'][0]['columns'][1]['types'] == ['int', 'str']
    _, entries = spreadsheet.select(record, sheet['id'], table_id='RulesTable', filters=[
        {'column_id': 'A', 'operator': 'eq', 'value': 'Reason'},
        {'column_id': 'B', 'operator': 'eq', 'value': '197'}])
    assert len(entries) == 1 and entries[0]['locator']['range'] == 'A2:C2'
    assert entries[0]['locator']['sheet'] == 'Rules'
    hidden = record['sheets'][1]
    assert hidden['visibility'] == 'hidden'
    assert hidden['cells']['A1']['formula'] == '=1+1'
    assert hidden['cells']['A1']['value'] is None
    assert hidden['merged_ranges'] == [{'range': 'B2:C2', 'anchor': 'B2'}]


def test_version_is_idempotent_and_sensitive_to_source_and_table_configuration():
    data = workbook(); first = spreadsheet.parse(data)
    assert spreadsheet.parse(data)['version'] == first['version']
    configured = [{'sheet': 'Rules', 'id': 'other', 'range': 'A1:C4', 'header_row': 1}]
    assert spreadsheet.parse(data, configured)['version'] != first['version']
    book = openpyxl.load_workbook(io.BytesIO(data)); book['Rules']['C2'] = 'Changed'
    stream = io.BytesIO(); book.save(stream)
    assert spreadsheet.parse(stream.getvalue())['version'] != first['version']


def test_unknown_headers_remain_raw_and_invalid_selectors_are_rejected():
    record = spreadsheet.parse(workbook()); sheet = record['sheets'][0]
    for values in [{'cell_range': 'A2:C2', 'table_id': 'RulesTable'}, {'cell_range': 'Z200:Z201'},
                   {'table_id': 'RulesTable', 'filters': [{'column_id': 'Z', 'operator': 'eq', 'value': 'x'}]}]:
        with pytest.raises(ValueError): spreadsheet.select(record, sheet['id'], **values)
    with pytest.raises(ValueError): spreadsheet.parse(b'%PDF-not-an-xlsx')
    book = openpyxl.Workbook(); book.active.append(['Some', 'Values'])
    stream = io.BytesIO(); book.save(stream)
    assert spreadsheet.parse(stream.getvalue())['sheets'][0]['tables'] == []


def test_large_cell_is_delivered_losslessly_and_cursor_cannot_change_binding():
    entries = [{'id': 'evidence', 'locator': {'kind': 'cells', 'sheet': 'Real', 'range': 'A1:A1'},
                'text': '业务方法😀' * 2000}]
    cursor = None; spans = []; texts = []
    for _ in range(50):
        result, delivered = read_batch({}, [{'page': 1, 'passages': entries}], identity='request',
            secret='secret', cursor=cursor, max_chars=2000)
        spans.extend(delivered); texts += [p['text'] for page in result['pages'] for p in page['passages']]
        cursor = result['reading']['next_cursor']
        if result['reading']['complete']: break
        with pytest.raises(ReadingCursorError):
            read_batch({}, [{'page': 1, 'passages': entries}], identity='different-invocation', secret='secret', cursor=cursor, max_chars=2000)
    assert ''.join(texts) == entries[0]['text']
    assert fully_read([(row[2], row[3]) for row in spans], len(entries[0]['text']))
    assert not fully_read([(spans[0][2], spans[0][3])], len(entries[0]['text']))
