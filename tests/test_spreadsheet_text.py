"""Text-index batches keep source record relationships; cells remain authoritative."""
import io
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.knowledge import spreadsheet
from app.parsing import spreadsheet_text


def record(long=False, configured=False):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Unseen Rules'
    sheet.append(['Kind', 'Code', 'Steps', 'Exceptions'])
    for row in range(2, 14):
        sheet.append(['Reason' if row % 2 else 'Status', 197, '步骤🙂 ' * (500 if long and row == 3 else 6), f'Exception {row}'])
    if not configured: sheet.add_table(Table(displayName='Methods', ref='A1:D13'))
    sheet['F17'] = 'Separate note, no inferred header'
    book.create_sheet('Another sheet')['C8'] = 'Independent fact'
    file = io.BytesIO(); book.save(file)
    config = [{'id': 'Configured', 'sheet': sheet.title, 'range': 'A1:D13', 'header_row': 1}] if configured else []
    return spreadsheet.parse(file.getvalue(), config)


def batches(text, meta):
    return [text[start:end] for start, end in meta['spreadsheet_text']['batches']]


def rows(piece):
    return [json.loads(line) for line in piece.splitlines() if line.startswith('{"cells"') or '"row_fragment"' in line]


def test_complete_records_keep_real_columns_coordinates_and_headers_in_every_batch():
    text, meta = spreadsheet_text.render(record(), limit=1100)
    parts = batches(text, meta); normal = []
    for part in parts:
        assert len(part) <= 1100
        assert 'worksheet=' in part and 'region=' in part and 'columns=' in part
        data = rows(part); normal += data
        for item in data:
            assert item['row_complete'] and 'row_fragment' not in item
            if item['range'].startswith('A'):
                assert '"header":"Kind"' in part and '"header":"Exceptions"' in part
                assert item['cells'][f'D{item["row"]}'] == f'Exception {item["row"]}'
                assert item['cells'][f'B{item["row"]}'] == 197
    assert sorted(r['row'] for r in normal if r['range'].startswith('A')) == list(range(2,14))
    assert any(r['range'] == 'F17:F17' for r in normal)
    assert any(r['range'] == 'C8:C8' for r in normal)
    assert ''.join(parts) == text


def test_only_oversized_row_splits_and_fragments_reconstruct_identical_record():
    text, meta = spreadsheet_text.render(record(long=True, configured=True), limit=1100)
    fragments = [r for p in batches(text, meta) for r in rows(p) if 'row_fragment' in r]
    assert len(fragments) > 1 and {r['row'] for r in fragments} == {3}
    assert all(not r['row_complete'] and r['range'] == 'A3:D3' for r in fragments)
    assert [r['continuation']['part'] for r in fragments] == list(range(1,len(fragments)+1))
    assert all(r['continuation']['parts'] == len(fragments) for r in fragments)
    original = json.loads(''.join(r['row_fragment'] for r in fragments))
    assert original['cells']['C3'] == '步骤🙂 ' * 500 and original['cells']['D3'] == 'Exception 3'
    assert all('Configured' in p and '"header":"Kind"' in p for p in batches(text,meta) if 'row_fragment' in p)


def test_build_uses_declared_batches_without_repartitioning_records(monkeypatch):
    from types import SimpleNamespace
    from app.services import pageindex_engine as engine
    text, meta = spreadsheet_text.render(record(), limit=1100)
    document = SimpleNamespace(file_type='xlsx', meta=meta)
    pieces = engine._document_pieces(document, text)
    assert [p for _,p in pieces] == batches(text,meta)
    import pytest
    with pytest.raises(ValueError): engine._document_pieces(document, text + 'changed')
