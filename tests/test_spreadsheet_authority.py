"""Native cells integrated with the real KB PostgreSQL authority in an isolated DB."""
import io
import time
from uuid import UUID

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableStyleInfo
from test_knowledge_evidence import isolated_database, headers
from app.core.database import SessionLocal
from app.models import User, KnowledgeBase, Document
from app.knowledge.access import request_headers
from app.knowledge.spreadsheet_service import SpreadsheetService, publish
from app.knowledge import store


@pytest.fixture
def cells(monkeypatch):
    from app.utils.minio_client import minio_client
    objects = {}
    monkeypatch.setattr(minio_client, 'get_object', lambda key: objects[key])
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Actual Rules'
    sheet.append(['Code Type', 'Code', 'Action']); sheet.append(['Reason', 197, 'Check prior authorization'])
    sheet.append(['Status', 197, 'A different workflow']); sheet.append(['Reason', '197', 'Text-code workflow'])
    sheet.add_table(Table(displayName='Rules', ref='A1:C4'))
    file = io.BytesIO(); book.save(file); objects['original.xlsx'] = file.getvalue()
    with SessionLocal() as db:
        owner = User(username='owner', email='owner@test.invalid', password_hash='fixture', role='user')
        outsider = User(username='outsider', email='outsider@test.invalid', password_hash='fixture', role='user')
        db.add_all([owner, outsider]); db.flush()
        kb = KnowledgeBase(owner_id=owner.id, name='Claim workflow knowledge', description='Applicable claim code methods')
        db.add(kb); db.flush()
        doc = Document(kb_id=kb.id, uploaded_by=owner.id, filename='original.xlsx', file_type='xlsx', minio_key='original.xlsx', status='indexed', full_text='Preserved original parsed text', meta={'index_type': 'embedding', 'previous_success': 'retained'})
        db.add(doc); db.commit(); ident = doc.id
        record = publish(doc, db); db.commit()
    return record, ident, objects


def call(name, values, actor='test-actor', invocation='invocation-1'):
    head = headers(actor, invocation); token = request_headers.set(head)
    try: return SpreadsheetService().call(name, values, head['x-carefold-interaction-scope'])
    finally: request_headers.reset(token)


def args(record): return {'document_id': record['document_id'], 'version': record['version']}


def test_discovery_and_original_cell_workflow_preserve_existing_index(cells, monkeypatch):
    record, ident, _ = cells
    # A native XLSX operation must not touch either PDF SDK entry.
    from app.knowledge import documents
    monkeypatch.setattr(documents, 'prepare_pdf', lambda *a: pytest.fail('Native cells used a PDF'))
    monkeypatch.setattr('pageindex.flash.page_index_flash', lambda *a, **kw: pytest.fail('Native cells used the PageIndex SDK'))
    collections = call('knowledge_discover', {})
    assert collections['collections'][0]['name'] == 'Claim workflow knowledge'
    catalog = call('knowledge_discover', {'kb_id': record['kb_id']})
    assert catalog['documents'][0]['read_methods'] == ['spreadsheet_read_cells']
    sheets = call('spreadsheet_get_structure', args(record)); sheet = sheets['sheets'][0]
    regions = call('spreadsheet_get_structure', args(record) | {'sheet_id': sheet['id']})
    columns = call('spreadsheet_get_structure', args(record) | {'sheet_id': sheet['id'], 'table_id': regions['regions'][0]['id']})
    assert [c['id'] for c in columns['columns']] == ['A', 'B', 'C']
    result = call('spreadsheet_read_cells', args(record) | {'sheet_id': sheet['id'], 'table_id': 'Rules', 'filters': [{'column_id': 'A', 'operator': 'eq', 'value': 'Reason'}, {'column_id': 'B', 'operator': 'eq', 'value': 197}]})
    assert len(result['evidence']) == 1
    item = result['evidence'][0]; assert item['locator']['range'] == 'A2:C2'
    assert 'pages' not in result and 'requested_pages' not in result['reading']
    monkeypatch.setattr('app.knowledge.models.support_check', lambda *a: {'supported': True, 'reason': 'Actual source supports it'})
    verified = call('spreadsheet_verify_citations', args(record) | {'citations': [{'sheet_id': sheet['id'], 'cell_range': 'A2:C2', 'evidence_id': item['id'], 'claim': 'The Reason code rule requires checking prior authorization.'}]})
    assert verified['citations'][0]['support_status'] == 'checked'
    assert verified['citations'][0]['locator']['kind'] == 'cells'
    assert 'page' not in verified['citations'][0]
    with SessionLocal() as db:
        doc = db.get(Document, ident)
        assert doc.full_text == 'Preserved original parsed text'
        assert doc.meta['index_type'] == 'embedding' and doc.meta['previous_success'] == 'retained'
        assert publish(doc, db)['version'] == record['version']; db.commit()


def test_native_acl_version_and_receipts_are_rechecked(cells, monkeypatch):
    record, ident, objects = cells; sheet = record['sheets'][0]
    selection = args(record) | {'sheet_id': sheet['id'], 'cell_range': 'A2:C2'}
    with pytest.raises(PermissionError): call('spreadsheet_read_cells', selection, actor='other-actor')
    read = call('spreadsheet_read_cells', selection)
    citation = args(record) | {'citations': [{'sheet_id': sheet['id'], 'cell_range': 'A2:C2', 'evidence_id': read['evidence'][0]['id'], 'claim': 'A claim'}]}
    with pytest.raises(PermissionError): call('spreadsheet_verify_citations', citation, invocation='new-invocation')
    monkeypatch.setattr('app.knowledge.models.support_check', lambda *a: {'supported': False, 'reason': 'Contradiction'})
    assert call('spreadsheet_verify_citations', citation)['citations'][0]['support_status'] == 'needs_review'
    objects['original.xlsx'] += b'changed'
    with pytest.raises(ValueError, match='changed'): call('spreadsheet_read_cells', selection)
    with SessionLocal() as db:
        doc = db.get(Document, ident); old = doc.meta['spreadsheet_evidence']
        objects['original.xlsx'] = b'invalid workbook'
        with pytest.raises(ValueError): publish(doc, db)
        db.rollback(); assert db.get(Document, ident).meta['spreadsheet_evidence'] == old
    objects['original.xlsx'] = b'denied source must not open'
    with SessionLocal() as db:
        kb = db.get(KnowledgeBase, UUID(record['kb_id'])); kb.owner_id = db.query(User).filter_by(username='outsider').one().id; db.commit()
    monkeypatch.setattr('app.knowledge.spreadsheet_service.minio_client.get_object', lambda *a: pytest.fail('Revoked source opened'))
    with pytest.raises(PermissionError): call('spreadsheet_read_cells', selection)


def test_partial_long_cells_cannot_be_verified_and_cursor_is_invocation_bound(cells, monkeypatch):
    record, ident, objects = cells
    book = openpyxl.load_workbook(io.BytesIO(objects['original.xlsx'])); book.active['C2'] = '长内容🙂' * 5000
    file = io.BytesIO(); book.save(file); objects['original.xlsx'] = file.getvalue()
    with SessionLocal() as db: record = publish(db.get(Document, ident), db); db.commit()
    sheet = record['sheets'][0]; selection = args(record) | {'sheet_id': sheet['id'], 'cell_range': 'C2:C2', 'max_chars': 3000}
    read = call('spreadsheet_read_cells', selection); first = read['evidence'][0]
    assert not read['reading']['complete']
    citation = args(record) | {'citations': [{'sheet_id': sheet['id'], 'cell_range': 'C2:C2', 'evidence_id': first['id'], 'claim': 'The cell is long.'}]}
    with pytest.raises(PermissionError): call('spreadsheet_verify_citations', citation)
    with pytest.raises(ValueError): call('spreadsheet_read_cells', selection | {'cursor': read['reading']['next_cursor']}, invocation='different')
    text = first['text']; rounds = 0
    while not read['reading']['complete']:
        read = call('spreadsheet_read_cells', selection | {'cursor': read['reading']['next_cursor']}); text += read['evidence'][0]['text']; rounds += 1
        assert len(__import__('json').dumps(read, ensure_ascii=False, separators=(',', ':')).encode('utf-16-le')) // 2 <= 3000
    assert text == __import__('app.knowledge.spreadsheet', fromlist=['select']).select(record, sheet['id'], 'C2:C2')[1][0]['text']
    assert rounds > 1
    monkeypatch.setattr('app.knowledge.models.support_check', lambda *a: {'supported': True, 'reason': 'Fully read'})
    assert call('spreadsheet_verify_citations', citation)['citations'][0]['support_status'] == 'checked'


def test_table_windows_preserve_all_matches_and_sparse_filtering(cells):
    record, ident, objects = cells
    book = openpyxl.load_workbook(io.BytesIO(objects['original.xlsx']))
    sheet = book.active
    for row in range(5, 605): sheet.append(['Reason', row, 'Workflow row'])
    sheet.tables['Rules'].ref = 'A1:C604'
    file = io.BytesIO(); book.save(file); objects['original.xlsx'] = file.getvalue()
    with SessionLocal() as db: record = publish(db.get(Document, ident), db); db.commit()
    selection = args(record) | {'sheet_id': record['sheets'][0]['id'], 'table_id': 'Rules', 'max_chars': 24000}
    seen = set(); read = None; cursors = set()
    for _ in range(40):
        read = call('spreadsheet_read_cells', selection | ({'cursor': read['reading']['next_cursor']} if read else {}))
        for entry in read['evidence']: seen.add(entry['locator']['range'])
        cursor = read['reading']['next_cursor']
        assert cursor is None or cursor not in cursors
        cursors.add(cursor)
        if read['reading']['complete']: break
    assert read['reading']['complete'] and len(seen) == 603
    rare = call('spreadsheet_read_cells', selection | {'filters': [{'column_id': 'B', 'operator': 'eq', 'value': 604}]})
    assert rare['reading']['complete'] and rare['evidence'][0]['locator']['range'] == 'A604:C604'


def test_catalog_respects_result_bound_and_keeps_paging():
    service = SpreadsheetService()
    items = [{'name': str(n), 'description': '内容🙂' * 1000} for n in range(50)]
    result = service.catalog(items, 50, 0, collections=items)
    assert result['next_offset'] == len(result['collections']) < 50
    assert len(__import__('json').dumps(result, ensure_ascii=False, separators=(',', ':')).encode('utf-16-le')) // 2 <= 64000


def test_admin_configuration_and_native_queue_failure_preserve_text_success(cells, monkeypatch):
    from app.api.documents import configure_spreadsheet
    from app.models import KbMember
    from app.workers import tasks
    from app.core.responses import APIError
    record, ident, _ = cells
    with SessionLocal() as db:
        owner = db.query(User).filter_by(username='owner').one()
        reader = db.query(User).filter_by(username='outsider').one()
        db.add(KbMember(kb_id=UUID(record['kb_id']), user_id=reader.id, role='viewer')); db.commit()
        with pytest.raises(APIError): configure_spreadsheet(UUID(record['kb_id']), ident, [], db, reader)
        with pytest.raises(APIError): configure_spreadsheet(UUID(record['kb_id']), ident, [{'sheet': 'Missing', 'id': 'other', 'range': 'A1:C4', 'header_row': 1}], db, owner)
        assert 'spreadsheet_tables' not in db.get(Document, ident).meta
    monkeypatch.setattr('app.parsing.parse_document', lambda *a: ('Successful text parse', {'parser': 'original-text-path'}))
    queued = []
    monkeypatch.setattr(tasks.chunk_index_task, 'delay', lambda *a: queued.append(a))
    def failed_queue(*a): raise RuntimeError('Native queue unavailable')
    monkeypatch.setattr(tasks.build_spreadsheet_evidence_task, 'delay', failed_queue)
    tasks.parse_document_task.run(str(ident), '00000000-0000-0000-0000-000000000001')
    with SessionLocal() as db:
        doc = db.get(Document, ident)
        assert doc.status == 'parsed' and doc.full_text == 'Successful text parse'
        assert doc.meta['spreadsheet_status'] == 'failed' and doc.meta['spreadsheet_evidence']['version'] == record['version']
        assert queued
